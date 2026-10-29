"""Synthetic test documents; no real or institutional information."""
import io
import unittest
from docx import Document
from engine import generate,preview_html

class DocxEngineTests(unittest.TestCase):
    def sample(self):
        doc=Document()
        p=doc.add_paragraph('Exchange: ')
        p.add_run('Original').bold=True
        p.add_run(' Company')
        doc.add_paragraph('Reference {{REFERENCE}}; issued {{DATE_LONG}}')
        table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='Jurisdiction'
        table.cell(0,1).text='{{INCORPORATION}}'
        doc.sections[0].header.paragraphs[0].text='Officer {{OFFICER}}'
        output=io.BytesIO();doc.save(output)
        return output.getvalue()

    def test_generate_replaces_across_runs_and_in_table_header(self):
        result,report=generate(self.sample(),{
            'reference':'EXAMPLE/42','date':'2026-10-29',
            'incorporation':'Example Jurisdiction','officer':'Example Officer',
            'rules':'[{"old":"Original Company","new":"Example Exchange"}]'
        })
        document=Document(io.BytesIO(result))
        self.assertIn('Example Exchange',document.paragraphs[0].text)
        self.assertIn('29th October 2026',document.paragraphs[1].text)
        self.assertIn('EXAMPLE/42',document.paragraphs[1].text)
        self.assertEqual(document.tables[0].cell(0,1).text,'Example Jurisdiction')
        self.assertIn('Example Officer',document.sections[0].header.paragraphs[0].text)
        self.assertGreater(len(report['replacements']),0)

    def test_unmatched_values_are_not_invented(self):
        result,report=generate(self.sample(),{'exchange':'Unknown','rules':'[]'})
        self.assertIn('Original Company',Document(io.BytesIO(result)).paragraphs[0].text)
        self.assertTrue(any('No placeholders' in w for w in report['warnings']))

    def test_html_preview_escapes_untrusted_content(self):
        result,_=generate(self.sample(),{'reference':'<script>alert(1)</script>'})
        preview=preview_html(result)
        self.assertIn('&lt;script&gt;',preview)
        self.assertNotIn('<script>',preview)

if __name__=='__main__':
    unittest.main()
