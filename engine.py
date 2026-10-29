"""Generic DOCX placeholder and manual contextual replacement engine."""
import io,json,html
from datetime import date
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

FIELDS=("exchange","entity","incorporation","principal","license","reference",
        "date","officer","designation","department","email","approver","subjects")

def iter_paragraphs(document):
    def visit(obj):
        for p in obj.paragraphs: yield p
        for table in obj.tables:
            for row in table.rows:
                for cell in row.cells: yield from visit(cell)
    yield from visit(document)
    for sec in document.sections:
        yield from visit(sec.header)
        yield from visit(sec.footer)

def replace_in_paragraph(p,old,new):
    if not old or old==new: return 0
    hits=0
    while old in p.text:
        runs=p.runs
        entire="".join(run.text for run in runs)
        i=entire.find(old)
        if i<0:break
        end=i+len(old)
        start_run=end_run=None
        cursor=0
        for index,run in enumerate(runs):
            next_cursor=cursor+len(run.text)
            if start_run is None and i<next_cursor: start_run=index;a=i-cursor
            if start_run is not None and end<=next_cursor:
                end_run=index;b=end-cursor;break
            cursor=next_cursor
        if start_run is None or end_run is None:break
        if start_run==end_run:
            r=runs[start_run];r.text=r.text[:a]+new+r.text[b:]
        else:
            tail=runs[end_run].text[b:]
            runs[start_run].text=runs[start_run].text[:a]+new
            for j in range(start_run+1,end_run):runs[j].text=""
            runs[end_run].text=tail
        hits+=1
        if hits>2000:raise ValueError("Too many replacements")
    return hits

def generate(original, fields):
    if len(original)>8*1024*1024:raise ValueError("DOCX exceeds 8 MB")
    doc=Document(io.BytesIO(original))
    pairs=[]
    for key in FIELDS:
        val=fields.get(key,"").strip()
        if val:
            pairs.extend([("{{"+key.upper()+"}}",val),("[["+key.upper()+"]]",val)])
    if fields.get("date"):
        d=date.fromisoformat(fields["date"])
        suffix='th' if 11<=d.day%100<=13 else {1:'st',2:'nd',3:'rd'}.get(d.day%10,'th')
        formatted=f"{d.day}{suffix} {d:%B %Y}"
        pairs.extend([("{{DATE_LONG}}",formatted),("[[DATE_LONG]]",formatted)])
    try:rules=json.loads(fields.get("rules","[]"))
    except json.JSONDecodeError:raise ValueError("Invalid replacement rules")
    if not isinstance(rules,list) or len(rules)>60:raise ValueError("Up to 60 rules are supported")
    for rule in rules:
        if not isinstance(rule,dict):raise ValueError("Invalid rule")
        old,new=rule.get("old"),rule.get("new")
        if not isinstance(old,str) or not isinstance(new,str) or len(old)>500 or len(new)>3000:
            raise ValueError("Rule strings invalid or too long")
        if old and old!=new:
            if old in new:raise ValueError("Replacement cannot contain its own search text")
            pairs.append((old,new))
    paragraphs=list(iter_paragraphs(doc))
    counters=[]
    for old,new in sorted(pairs,key=lambda item:len(item[0]),reverse=True):
        hits=sum(replace_in_paragraph(p,old,new) for p in paragraphs)
        if hits:counters.append({"text":old,"occurrences":hits})
    output=io.BytesIO();doc.save(output)
    warnings=["Review all edited paragraphs, tables, headers, footers, signatures and legal statements.",
              "The application cannot independently verify corporate or licensing details."]
    if not counters:warnings.insert(0,"No placeholders or rules matched the imported document.")
    return output.getvalue(),{"replacements":counters,"warnings":warnings}

def preview_html(docx_bytes):
    """Escaped, simplified view if local Microsoft Word rendering unavailable."""
    d=Document(io.BytesIO(docx_bytes))
    fragments=[]
    for element in d.element.body.iterchildren():
        if element.tag==qn("w:p"):
            p=Paragraph(element,d)
            fragments.append("<p>"+html.escape(p.text or " ")+"</p>")
        elif element.tag==qn("w:tbl"):
            t=Table(element,d)
            fragments.append("<table>")
            for row in t.rows:
                fragments.append("<tr>"+"".join("<td>"+html.escape(cell.text)+"</td>" for cell in row.cells)+"</tr>")
            fragments.append("</table>")
    return '<article class="document">'+''.join(fragments)+"</article>"
